import io

from openpyxl import load_workbook

from stockroom.extensions import db
from stockroom.models import Order, Project, Role, User


def count(model):
    return db.session.scalar(db.select(db.func.count()).select_from(model))


def test_pm_cannot_access_admin(client, login, data):
    login(data["pm_a"])
    assert client.get("/admin/").status_code == 403
    assert client.post("/admin/projects", data={"name": "X"}).status_code == 403


def test_dashboard_renders(client, login, data):
    login(data["admin"])
    resp = client.get("/admin/")
    assert resp.status_code == 200
    assert b"Project A" in resp.data and b"Pat A" in resp.data


def test_create_project_and_add_pm(client, login, data):
    login(data["admin"])
    client.post("/admin/projects", data={"name": "Rover"})
    project = db.session.scalar(db.select(Project).filter_by(name="Rover"))
    client.post(
        "/admin/members",
        data={"name": "Sam", "email": "Sam@Example.com ", "project_id": project.id},
    )
    sam = db.session.scalar(db.select(User).filter_by(email="sam@example.com"))
    assert sam.role is Role.PM and sam.project_id == project.id


def test_duplicate_email_rejected(client, login, data):
    login(data["admin"])
    resp = client.post(
        "/admin/members",
        data={"name": "Dup", "email": "pm.b@example.com", "project_id": data["a"]},
        follow_redirects=True,
    )
    assert b"already has access" in resp.data


def test_delete_project_cascades(client, login, data):
    login(data["admin"])
    client.post(f"/admin/projects/{data['a']}/delete")
    db.session.expire_all()
    assert db.session.get(Project, data["a"]) is None
    assert db.session.get(User, data["pm_a"]) is None
    assert count(Order) == 0


def test_only_superuser_manages_admins(client, login, data):
    other = User(email="other@example.com", name="Other", role=Role.ADMIN)
    db.session.add(other)
    db.session.commit()
    login(data["admin"])
    assert (
        client.post("/admin/admins", data={"name": "X", "email": "x@example.com"}).status_code
        == 403
    )
    assert client.post(f"/admin/users/{other.id}/delete").status_code == 403
    assert client.post("/admin/reset", data={"confirm": "RESET"}).status_code == 403


def test_superuser_resets_data_keeping_admins(client, login, data):
    boss = User(email="boss@example.com", name="Boss", role=Role.ADMIN, is_superuser=True)
    db.session.add(boss)
    db.session.commit()
    login(boss.id)
    client.post("/admin/reset", data={"confirm": "nope"})
    assert count(Project) == 2
    client.post("/admin/reset", data={"confirm": "RESET"})
    assert count(Project) == 0 and count(Order) == 0
    assert {u.role for u in db.session.scalars(db.select(User))} == {Role.ADMIN}


def test_export_xlsx(client, login, data):
    login(data["admin"])
    resp = client.post("/admin/export", data={"all": "1"})
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.data))
    assert wb.sheetnames == ["Project A", "Project B"]
    rows = list(wb["Project A"].values)
    assert rows[1][9] == "Resistor"
    assert rows[1][-1] == 1.0


def test_export_selected_projects(client, login, data):
    login(data["admin"])
    resp = client.post("/admin/export", data={"project_id": [str(data["b"])]})
    assert load_workbook(io.BytesIO(resp.data)).sheetnames == ["Project B"]
